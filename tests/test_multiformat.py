import json
from pathlib import Path

import pytest
from docx import Document
from PIL import Image
from pptx import Presentation
from pptx.util import Inches

from ddr_rag.chunker import (
    ChunkingError,
    _build_hybrid_chunker,
    _load_docling_document,
    build_chunk_records,
)
from ddr_rag.config import load_settings
from ddr_rag.parser import _build_converter, parse_document, parsed_output_paths
from ddr_rag.schemas import DocumentRecord
from test_config import write_config


def make_record(file: Path, *, doc_id: str, document_type: str, authority: str) -> DocumentRecord:
    return DocumentRecord(
        doc_id=doc_id,
        title=f"{doc_id} source",
        revision="R1",
        vendor="Test Engineering",
        file=file,
        document_type=document_type,
        language="zh-CN",
        memory_types=["DDR4"],
        applicable_parts=["RK3568"],
        authority=authority,
        authority_score=40 if authority == "unverified_note" else 90,
        status="active",
        confidentiality="internal",
    )


def test_docx_parse_and_chunk_uses_structure_location_without_fake_pages(tmp_path: Path) -> None:
    config_path = tmp_path / "config.yaml"
    write_config(config_path)
    settings = load_settings(config_path)
    source = tmp_path / "data/raw/internal/rk3568-note.docx"
    source.parent.mkdir(parents=True)
    word = Document()
    word.add_heading("问题分析", level=1)
    paragraph = word.add_paragraph("RK3568 DDR4 时钟在本项目中按差分阻抗控制。")
    word.add_comment(paragraph.runs, text="这条 Word 批注不应入库。", author="Tester", initials="T")
    table = word.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "项目"
    table.cell(0, 1).text = "经验值"
    table.cell(1, 0).text = "CLK"
    table.cell(1, 1).text = "按原厂约束复核"
    word.save(source)
    record = make_record(
        Path("data/raw/internal/rk3568-note.docx"),
        doc_id="RK3568_NOTE",
        document_type="engineering_experience",
        authority="unverified_note",
    )

    report = parse_document(settings, record, _build_converter(settings))
    parsed = _load_docling_document(parsed_output_paths(settings, record)[1])
    records = build_chunk_records(
        record,
        parsed,
        _build_hybrid_chunker(settings),
        settings.chunking.max_tokens,
        source_file_hash="a" * 64,
        parser_version="2.107.0",
    )

    assert report["status"] == "success"
    assert report["source_format"] == "docx"
    assert "comments" in report["excluded_scope"]
    assert records
    assert all(chunk.source_format == "docx" for chunk in records)
    assert all(not chunk.page_numbers and chunk.page_start is None for chunk in records)
    assert all(chunk.source_location.kind == "docx_elements" for chunk in records)
    assert all(chunk.source_location.element_refs for chunk in records)
    assert all(chunk.source_location.excerpt for chunk in records)
    assert all(chunk.document_type == "engineering_experience" for chunk in records)
    assert all("批注不应入库" not in chunk.text for chunk in records)
    markdown = parsed_output_paths(settings, record)[0].read_text(encoding="utf-8")
    assert "source_format: docx" in markdown
    assert "RK3568 DDR4" in markdown


def test_pptx_keeps_original_slide_numbers_and_excludes_hidden_and_blank_slides(tmp_path: Path) -> None:
    config_path = tmp_path / "config.yaml"
    write_config(config_path)
    settings = load_settings(config_path)
    source = tmp_path / "data/raw/internal/rk3568-review.pptx"
    source.parent.mkdir(parents=True)
    presentation = Presentation()
    visible = presentation.slides.add_slide(presentation.slide_layouts[1])
    visible.shapes.title.text = "DDR4 布线复盘"
    visible.placeholders[1].text = "RK3568 CLK 规则需要回查原厂资料。"
    visible.notes_slide.notes_text_frame.text = "这条 PowerPoint 备注不应入库。"
    hidden = presentation.slides.add_slide(presentation.slide_layouts[1])
    hidden.shapes.title.text = "内部隐藏页"
    hidden.placeholders[1].text = "这段隐藏内容不应入库。"
    hidden._element.set("show", "0")
    presentation.slides.add_slide(presentation.slide_layouts[6])
    last = presentation.slides.add_slide(presentation.slide_layouts[6])
    textbox = last.shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(1))
    textbox.text_frame.text = "第 4 张可见幻灯片的 DDR4 复核结论。"
    presentation.save(source)
    record = make_record(
        Path("data/raw/internal/rk3568-review.pptx"),
        doc_id="RK3568_REVIEW",
        document_type="engineering_experience",
        authority="approved_experience",
    )

    report = parse_document(settings, record, _build_converter(settings))
    parsed = _load_docling_document(parsed_output_paths(settings, record)[1])
    records = build_chunk_records(
        record,
        parsed,
        _build_hybrid_chunker(settings, merge_peers=False),
        settings.chunking.max_tokens,
        hidden_slide_numbers=set(report["hidden_slide_numbers"]),
        source_file_hash="b" * 64,
        parser_version="2.107.0",
    )

    assert report["source_format"] == "pptx"
    assert report["source_page_count"] == 4
    assert report["hidden_slide_numbers"] == [2]
    slide_numbers = {
        number
        for chunk in records
        for number in chunk.source_location.slide_numbers
    }
    assert slide_numbers == {1, 4}
    assert all(chunk.source_format == "pptx" for chunk in records)
    assert all(not chunk.page_numbers and chunk.page_start is None for chunk in records)
    assert all(len(chunk.source_location.slide_numbers) == 1 for chunk in records)
    assert all("隐藏内容" not in chunk.text for chunk in records)
    assert all("备注不应入库" not in chunk.text for chunk in records)


def test_picture_only_pptx_is_not_treated_as_complete_text_evidence(tmp_path: Path) -> None:
    config_path = tmp_path / "config.yaml"
    write_config(config_path)
    settings = load_settings(config_path)
    source = tmp_path / "data/raw/internal/picture-only.pptx"
    source.parent.mkdir(parents=True)
    picture = tmp_path / "topology.png"
    Image.new("RGB", (20, 20), "white").save(picture)
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    slide.shapes.add_picture(str(picture), Inches(1), Inches(1))
    presentation.save(source)
    record = make_record(
        Path("data/raw/internal/picture-only.pptx"),
        doc_id="PICTURE_ONLY",
        document_type="engineering_experience",
        authority="unverified_note",
    )

    report = parse_document(settings, record, _build_converter(settings))
    parsed = _load_docling_document(parsed_output_paths(settings, record)[1])

    assert report["picture_count"] == 1
    with pytest.raises(ChunkingError, match="produced no chunks"):
        build_chunk_records(
            record,
            parsed,
            _build_hybrid_chunker(settings, merge_peers=False),
            settings.chunking.max_tokens,
        )


def test_corrupt_docx_records_failure_log(tmp_path: Path) -> None:
    config_path = tmp_path / "config.yaml"
    write_config(config_path)
    settings = load_settings(config_path)
    source = tmp_path / "data/raw/internal/corrupt.docx"
    source.parent.mkdir(parents=True)
    source.write_bytes(b"not an Office Open XML document")
    record = make_record(
        Path("data/raw/internal/corrupt.docx"),
        doc_id="CORRUPT_DOCX",
        document_type="engineering_experience",
        authority="unverified_note",
    )

    with pytest.raises(Exception):
        parse_document(settings, record, _build_converter(settings))

    log_path = settings.resolve_path(settings.paths.logs) / "ingestion/CORRUPT_DOCX.parse.json"
    payload = json.loads(log_path.read_text(encoding="utf-8"))
    assert payload["status"] == "failure"
    assert payload["source_format"] == "docx"
    assert payload["error_type"]
