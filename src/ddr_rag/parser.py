"""Multi-format Docling conversion with deterministic outputs and audit logs."""

from __future__ import annotations

import json
import os
import time
from collections.abc import Callable
from datetime import datetime, timezone
from importlib.metadata import version
from pathlib import Path
from typing import Any

from ddr_rag.config import AppSettings
from ddr_rag.schemas import DocumentRecord


class ParseError(RuntimeError):
    pass


def parsed_output_paths(settings: AppSettings, document: DocumentRecord) -> tuple[Path, Path]:
    parsed_root = settings.resolve_path(settings.paths.parsed)
    return (
        parsed_root / "Markdown" / f"{document.doc_id}.md",
        parsed_root / "JSON" / f"{document.doc_id}.json",
    )


def _build_converter(settings: AppSettings) -> Any:
    from docling.datamodel.accelerator_options import AcceleratorOptions
    from docling.datamodel.base_models import InputFormat
    from docling.datamodel.pipeline_options import PdfPipelineOptions
    from docling.document_converter import DocumentConverter, PdfFormatOption

    options = PdfPipelineOptions(
        do_ocr=settings.parser.enable_ocr,
        do_table_structure=settings.parser.enable_table_structure,
        accelerator_options=AcceleratorOptions(
            device=settings.parser.device,
            num_threads=settings.parser.num_threads,
        ),
        document_timeout=float(settings.parser.document_timeout_seconds),
        ocr_batch_size=1,
        layout_batch_size=1,
        table_batch_size=1,
        queue_max_size=8,
        generate_page_images=False,
        generate_picture_images=False,
    )
    return DocumentConverter(
        allowed_formats=[InputFormat.PDF, InputFormat.DOCX, InputFormat.PPTX],
        format_options={InputFormat.PDF: PdfFormatOption(pipeline_options=options)},
    )


def _pptx_hidden_slide_numbers(source: Path) -> list[int]:
    """Return original one-based slide numbers that PowerPoint marks hidden."""

    from pptx import Presentation

    presentation = Presentation(str(source))
    return [
        index
        for index, slide in enumerate(presentation.slides, start=1)
        if str(slide._element.get("show", "1")).lower() in {"0", "false", "off", "no"}
    ]


def _error_payload(error: Any) -> dict[str, Any]:
    if hasattr(error, "model_dump"):
        return error.model_dump(mode="json")
    return {"message": str(error)}


def _write_log(settings: AppSettings, doc_id: str, payload: dict[str, Any]) -> Path:
    log_dir = settings.resolve_path(settings.paths.logs) / "ingestion"
    log_dir.mkdir(parents=True, exist_ok=True)
    log_path = log_dir / f"{doc_id}.parse.json"
    log_path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    return log_path


def parse_document(
    settings: AppSettings,
    document: DocumentRecord,
    converter: Any,
    replace: bool = False,
    progress: Callable[[str], None] | None = None,
) -> dict[str, Any]:
    """Parse one registered PDF/DOCX/PPTX and export Markdown plus Docling JSON."""

    from docling.datamodel.base_models import ConversionStatus
    from docling_core.types.doc import DoclingDocument
    source = settings.resolve_path(document.file)
    source_format = document.source_format or source.suffix.lower().lstrip(".")
    markdown_path, json_path = parsed_output_paths(settings, document)
    required_outputs = []
    if settings.parser.export_markdown:
        required_outputs.append(markdown_path)
    if settings.parser.export_json:
        required_outputs.append(json_path)

    if required_outputs and all(path.is_file() for path in required_outputs) and not replace:
        payload = {
            "doc_id": document.doc_id,
            "status": "skipped",
            "reason": "outputs_exist",
            "source": str(source),
            "completed_at": datetime.now(timezone.utc).isoformat(),
        }
        _write_log(settings, document.doc_id, payload)
        return payload

    started = time.perf_counter()
    base_payload: dict[str, Any] = {
        "doc_id": document.doc_id,
        "source": str(source),
        "source_format": source_format,
        "parser": "docling",
        "parser_version": version("docling"),
        "ocr_enabled": settings.parser.enable_ocr,
        "table_structure_enabled": settings.parser.enable_table_structure,
        "started_at": datetime.now(timezone.utc).isoformat(),
    }

    try:
        accepted = {ConversionStatus.SUCCESS, ConversionStatus.PARTIAL_SUCCESS}
        batch_size = settings.parser.page_batch_size if source_format == "pdf" else None
        chunk_documents: list[Any] = []
        conversion_errors: list[dict[str, Any]] = []
        statuses: list[Any] = []
        batches: list[dict[str, Any]] = []
        source_page_count: int | None = None
        hidden_slide_numbers: list[int] = []

        if source_format == "pdf":
            from pypdf import PdfReader

            source_page_count = len(PdfReader(source).pages)
            ranges: list[tuple[int, int] | None] = [
                (page_start, min(page_start + settings.parser.page_batch_size - 1, source_page_count))
                for page_start in range(1, source_page_count + 1, settings.parser.page_batch_size)
            ]
        else:
            ranges = [None]
            if source_format == "pptx":
                from pptx import Presentation

                source_page_count = len(Presentation(str(source)).slides)
                hidden_slide_numbers = _pptx_hidden_slide_numbers(source)

        for source_range in ranges:
            batch_started = time.perf_counter()
            convert_kwargs: dict[str, Any] = {"raises_on_error": False}
            if source_range is not None:
                page_start, page_end = source_range
                convert_kwargs["page_range"] = source_range
                if progress:
                    progress(
                        f"{document.doc_id}: parsing pages {page_start}-{page_end} "
                        f"of {source_page_count}"
                    )
            elif progress:
                progress(f"{document.doc_id}: parsing {source_format.upper()} document")
            result = converter.convert(
                source,
                **convert_kwargs,
            )
            if result.status not in accepted:
                details = [_error_payload(error) for error in result.errors]
                location = (
                    f"pages {source_range[0]}-{source_range[1]}"
                    if source_range is not None
                    else source_format.upper()
                )
                raise ParseError(
                    f"Docling conversion failed for {location} "
                    f"with status={result.status.value}: {details}"
                )
            chunk_documents.append(result.document)
            statuses.append(result.status)
            conversion_errors.extend(_error_payload(error) for error in result.errors)
            batch_record: dict[str, Any] = {
                "status": result.status.value,
                "duration_seconds": round(time.perf_counter() - batch_started, 3),
            }
            if source_range is not None:
                batch_record.update(page_start=source_range[0], page_end=source_range[1])
            batches.append(batch_record)

        combined_document = (
            chunk_documents[0]
            if len(chunk_documents) == 1
            else DoclingDocument.concatenate(chunk_documents)
        )
        overall_status = (
            ConversionStatus.PARTIAL_SUCCESS
            if ConversionStatus.PARTIAL_SUCCESS in statuses
            else ConversionStatus.SUCCESS
        )

        markdown_path.parent.mkdir(parents=True, exist_ok=True)
        json_path.parent.mkdir(parents=True, exist_ok=True)
        temp_markdown = markdown_path.with_suffix(".md.tmp")
        temp_json = json_path.with_suffix(".json.tmp")

        if settings.parser.export_markdown:
            metadata = (
                f"<!-- doc_id: {document.doc_id} -->\n"
                f"<!-- source_file: {document.file.as_posix()} -->\n"
                f"<!-- source_format: {source_format} -->\n"
                f"<!-- revision: {document.revision} -->\n\n"
            )
            markdown = combined_document.export_to_markdown(
                page_break_placeholder="\n\n<!-- page-break -->\n\n",
                image_placeholder="<!-- figure: refer to original source file -->",
            )
            temp_markdown.write_text(metadata + markdown, encoding="utf-8")

        if settings.parser.export_json:
            combined_document.save_as_json(
                temp_json,
                indent=2,
                coord_precision=2,
                confid_precision=3,
            )

        if settings.parser.export_markdown:
            os.replace(temp_markdown, markdown_path)
        if settings.parser.export_json:
            os.replace(temp_json, json_path)

        payload = {
            **base_payload,
            "status": overall_status.value,
            "completed_at": datetime.now(timezone.utc).isoformat(),
            "duration_seconds": round(time.perf_counter() - started, 3),
            "page_count": len(combined_document.pages),
            "source_page_count": source_page_count,
            "text_item_count": len(combined_document.texts),
            "table_count": len(combined_document.tables),
            "picture_count": len(combined_document.pictures),
            "page_batch_size": batch_size,
            "hidden_slide_numbers": hidden_slide_numbers,
            "excluded_scope": (
                ["comments", "notes", "hidden_slides"]
                if source_format == "pptx"
                else (["comments"] if source_format == "docx" else [])
            ),
            "batches": batches,
            "markdown": str(markdown_path) if settings.parser.export_markdown else None,
            "json": str(json_path) if settings.parser.export_json else None,
            "errors": conversion_errors,
        }
        _write_log(settings, document.doc_id, payload)
        return payload
    except Exception as exc:
        payload = {
            **base_payload,
            "status": "failure",
            "completed_at": datetime.now(timezone.utc).isoformat(),
            "duration_seconds": round(time.perf_counter() - started, 3),
            "error_type": type(exc).__name__,
            "error": str(exc),
        }
        _write_log(settings, document.doc_id, payload)
        raise


def parse_documents(
    settings: AppSettings,
    documents: list[DocumentRecord],
    replace: bool = False,
    progress: Callable[[str], None] | None = None,
) -> list[dict[str, Any]]:
    if settings.parser.engine.lower() != "docling":
        raise ParseError(f"Unsupported parser engine: {settings.parser.engine}")
    converter = _build_converter(settings)
    reports: list[dict[str, Any]] = []
    for document in documents:
        reports.append(
            parse_document(
                settings,
                document,
                converter,
                replace=replace,
                progress=progress,
            )
        )
    return reports
